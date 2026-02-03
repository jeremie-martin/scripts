#!/usr/bin/env python3
"""
jamatrace_analyze.py

Offline analyzer that:
  - reads a Jama graph DB JSON produced by jamagraph_export.py
  - reads a traceability config JSON (same format you already use)
  - classifies each node (type_key) from config's item_definitions
  - evaluates boolean link rules (upstream/downstream) from linking_rules
  - prints a human-readable report to stdout
  - writes an Excel workbook with results and a missing-links sheet

NO live Jama calls here.

Usage:
  python jamatrace_analyze.py --db-file jama_graph.json --config traceability_config.json \
    [--output-xlsx traceability_report.xlsx] [--keys-file subset.txt] [--report-format text|json|csv] [--verbose]
"""

from __future__ import annotations

import argparse
import csv
import io
import json
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


# ------------------------------- Data Models -------------------------------


@dataclass
class Node:
    id: int
    document_key: str
    name: str
    project_id: Optional[int]
    description_html: str
    url: str
    item_type_id: Optional[int]
    item_type_label: Optional[str]
    type_key: Optional[str] = None  # classification from config


@dataclass
class CheckResult:
    direction: str  # "upstream" | "downstream"
    expression: str
    passed: bool
    missing_types: List[str]
    found_nodes: List[Node]


@dataclass
class ItemResult:
    node: Node
    overall_passed: bool
    checks: List[CheckResult]


# --------------------------- DB / Config loading ---------------------------


def load_db(db_path: Path) -> Tuple[Dict[int, Node], List[Tuple[int, int]], List[str]]:
    data = json.loads(db_path.read_text(encoding="utf-8"))
    nodes_json = data.get("nodes", {})
    edges_json = data.get("edges", [])
    seeds = data.get("seed_keys", [])

    nodes: Dict[int, Node] = {}
    for k, v in nodes_json.items():
        nid = int(v["id"])
        nodes[nid] = Node(
            id=nid,
            document_key=v.get("document_key", ""),
            name=v.get("name", ""),
            project_id=v.get("project_id"),
            description_html=v.get("description_html", ""),
            url=v.get("url", ""),
            item_type_id=v.get("item_type_id"),
            item_type_label=v.get("item_type_label"),
            type_key=v.get("type_key"),
        )

    edges: List[Tuple[int, int]] = []
    for e in edges_json:
        edges.append((int(e["from"]), int(e["to"])))

    return nodes, edges, seeds


def load_config(cfg_path: Path) -> dict:
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    for section in ("item_definitions", "linking_rules"):
        if section not in cfg:
            raise ValueError(f"Missing required section in config: {section}")
    return cfg


# ---------------------------- Classification ----------------------------


def classify_nodes(nodes: Dict[int, Node], config: dict):
    defs = config.get("item_definitions", {})
    for n in nodes.values():
        # if type_key already present in DB, keep it; otherwise infer
        if n.type_key:
            continue
        inferred = None
        for type_key, definition in defs.items():
            if _matches_definition(n, definition):
                inferred = type_key
                break
        n.type_key = inferred


def _matches_definition(node: Node, definition) -> bool:
    if isinstance(definition, str):
        return node.document_key.startswith(definition)

    if not isinstance(definition, dict):
        return False

    pattern = definition.get("pattern", "")
    if pattern and not node.document_key.startswith(pattern):
        return False

    name = node.name or ""
    if "name_contains" in definition:
        if definition["name_contains"] not in name:
            return False

    if "name_contains_any" in definition:
        if not any(tok in name for tok in definition["name_contains_any"]):
            return False

    return True


# --------------------------- Relationship views ---------------------------


def build_index(edges: List[Tuple[int, int]]) -> Tuple[Dict[int, List[int]], Dict[int, List[int]]]:
    """Return (downstream_index, upstream_index)."""
    down: Dict[int, List[int]] = defaultdict(list)  # node -> [to_ids]
    up: Dict[int, List[int]] = defaultdict(list)  # node -> [from_ids]
    for fr, to in edges:
        down[fr].append(to)
        up[to].append(fr)
    return down, up


# ------------------------------ Rule engine ------------------------------


def eval_boolean_expression(expression: str, related_nodes: List[Node], config: dict) -> Tuple[bool, List[str], List[Node]]:
    """
    Replace type tokens with True/False based on presence in related_nodes.
    Supports "AND" / "OR". Parentheses OK.
    """
    expr = (expression or "").strip()
    if not expr:
        return True, [], []

    defs = config.get("item_definitions", {})
    type_keys = {k for k in defs.keys() if k in expr}

    by_type: Dict[str, List[Node]] = defaultdict(list)
    for n in related_nodes:
        if n.type_key:
            by_type[n.type_key].append(n)

    # Map each type token to True/False
    type_exists = {}
    found: List[Node] = []
    for tk in type_keys:
        present = len(by_type.get(tk, [])) > 0
        type_exists[tk] = present
        if present:
            found.extend(by_type[tk])

    eval_expr = expr
    for tk in sorted(type_keys, key=len, reverse=True):
        eval_expr = eval_expr.replace(tk, "True" if type_exists.get(tk, False) else "False")
    eval_expr = eval_expr.replace(" AND ", " and ").replace(" OR ", " or ")

    try:
        result = bool(eval(eval_expr, {"__builtins__": {}}, {}))
    except Exception:
        return False, sorted(type_keys), found

    missing = []
    if not result:
        missing = sorted([tk for tk in type_keys if not type_exists.get(tk, False)])

    return result, missing, found


def analyze_traceability(
    nodes: Dict[int, Node],
    edges: List[Tuple[int, int]],
    config: dict,
    restrict_doc_keys: Optional[Set[str]] = None,
) -> Tuple[List[ItemResult], dict, List[Tuple[Node, str, str]]]:
    """
    Returns:
      - list of ItemResult
      - summary stats dict
      - missing_links list of tuples (node, direction, missing_type)
    """
    down, up = build_index(edges)
    results: List[ItemResult] = []
    summary = {
        "total_items": 0,
        "items_evaluated": 0,
        "items_passed": 0,
        "items_failed": 0,
        "total_checks": 0,
        "checks_passed": 0,
        "checks_failed": 0,
    }
    missing_links: List[Tuple[Node, str, str]] = []

    defs = config.get("item_definitions", {})
    rules = config.get("linking_rules", {})

    # candidates = nodes filtered by optional restriction
    candidate_nodes = [n for n in nodes.values() if (not restrict_doc_keys or n.document_key in restrict_doc_keys)]
    summary["total_items"] = len(candidate_nodes)

    for n in candidate_nodes:
        tkey = n.type_key
        rule_block = rules.get(tkey, {}) if tkey else {}

        checks: List[CheckResult] = []
        # Build related sets
        upstream_nodes = [nodes[pid] for pid in up.get(n.id, []) if pid in nodes]
        downstream_nodes = [nodes[cid] for cid in down.get(n.id, []) if cid in nodes]

        # Evaluate upstream
        if "upstream_required" in rule_block:
            passed, missing, found = eval_boolean_expression(rule_block["upstream_required"], upstream_nodes, config)
            checks.append(CheckResult("upstream", rule_block["upstream_required"], passed, missing, found))
            summary["total_checks"] += 1
            if passed:
                summary["checks_passed"] += 1
            else:
                summary["checks_failed"] += 1
                for m in missing:
                    missing_links.append((n, "upstream", m))

        # Evaluate downstream
        if "downstream_required" in rule_block:
            passed, missing, found = eval_boolean_expression(rule_block["downstream_required"], downstream_nodes, config)
            checks.append(CheckResult("downstream", rule_block["downstream_required"], passed, missing, found))
            summary["total_checks"] += 1
            if passed:
                summary["checks_passed"] += 1
            else:
                summary["checks_failed"] += 1
                for m in missing:
                    missing_links.append((n, "downstream", m))

        # Only count as evaluated if there was at least one check
        if checks:
            summary["items_evaluated"] += 1
            overall = all(c.passed for c in checks)
            if overall:
                summary["items_passed"] += 1
            else:
                summary["items_failed"] += 1

            results.append(ItemResult(node=n, overall_passed=overall, checks=checks))

    return results, summary, missing_links


# ------------------------------ Reporting ------------------------------


def type_desc(config: dict, type_key: Optional[str]) -> str:
    if not type_key:
        return "Unknown"
    d = config.get("item_definitions", {}).get(type_key)
    if isinstance(d, dict):
        return d.get("description", type_key)
    return type_key


def render_stdout(results: List[ItemResult], summary: dict, config: dict) -> str:
    out: List[str] = []
    out.append("=" * 80)
    out.append("JAMA TRACEABILITY REPORT (Offline DB)")
    out.append("=" * 80)

    grouped: Dict[str, List[ItemResult]] = defaultdict(list)
    for r in results:
        grouped[type_desc(config, r.node.type_key)].append(r)

    for tdesc, group in grouped.items():
        out.append(f"\n--- {tdesc} ---")
        for r in group:
            status = "✓ PASS" if r.overall_passed else "✗ FAIL"
            out.append(f"\n{status} {r.node.document_key}")
            out.append(f"    {r.node.name}")
            for c in r.checks:
                pretty = c.expression
                for tk, td in config.get("item_definitions", {}).items():
                    if isinstance(td, dict) and "description" in td:
                        pretty = pretty.replace(tk, td["description"])
                prefix = "✓" if c.passed else "✗"
                out.append(f"    {prefix} {c.direction.title()}: {pretty}")
                if c.found_nodes:
                    found = ", ".join(f"{n.document_key} ({type_desc(config, n.type_key)})" for n in c.found_nodes)
                    out.append(f"      Found: {found}")
                if not c.passed and c.missing_types:
                    out.append(f"      Missing: {', '.join(type_desc(config, m) for m in c.missing_types)}")

    out.append("\n" + "=" * 80)
    out.append("SUMMARY")
    out.append("=" * 80)
    out.append(f"Items in scope:   {summary['total_items']}")
    out.append(f"Items evaluated:  {summary['items_evaluated']}")
    out.append(f"Items passed:     {summary['items_passed']}")
    out.append(f"Items failed:     {summary['items_failed']}")
    out.append(f"Checks total:     {summary['total_checks']}")
    out.append(f"Checks passed:    {summary['checks_passed']}")
    out.append(f"Checks failed:    {summary['checks_failed']}")
    return "\n".join(out)


def render_structured(results: List[ItemResult], summary: dict, config: dict, fmt: str) -> str:
    if fmt == "json":
        payload = {
            "summary": summary,
            "results": [
                {
                    "item": {
                        "id": r.node.id,
                        "document_key": r.node.document_key,
                        "name": r.node.name,
                        "type": r.node.type_key,
                        "type_description": type_desc(config, r.node.type_key),
                        "url": r.node.url,
                    },
                    "overall_passed": r.overall_passed,
                    "checks": [
                        {
                            "direction": c.direction,
                            "expression": c.expression,
                            "passed": c.passed,
                            "missing_types": c.missing_types,
                            "found": [
                                {
                                    "id": n.id,
                                    "document_key": n.document_key,
                                    "name": n.name,
                                    "type": n.type_key,
                                    "type_description": type_desc(config, n.type_key),
                                    "url": n.url,
                                }
                                for n in c.found_nodes
                            ],
                        }
                        for c in r.checks
                    ],
                }
                for r in results
            ],
        }
        return json.dumps(payload, indent=2)

    # csv
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(
        [
            "Document Key",
            "Name",
            "Type",
            "Type Description",
            "Overall",
            "Direction",
            "Expression",
            "Check Passed",
            "Found Links",
            "Missing Types",
        ]
    )
    for r in results:
        base = [
            r.node.document_key,
            r.node.name,
            r.node.type_key or "",
            type_desc(config, r.node.type_key),
            "PASS" if r.overall_passed else "FAIL",
        ]
        if not r.checks:
            w.writerow(base + ["", "", "", "", ""])
        else:
            for c in r.checks:
                found = "; ".join(f"{n.document_key} ({type_desc(config, n.type_key)})" for n in c.found_nodes)
                missing = "; ".join(type_desc(config, m) for m in c.missing_types)
                w.writerow(base + [c.direction, c.expression, "PASS" if c.passed else "FAIL", found, missing])
    return buf.getvalue()


# ------------------------------ Excel writer ------------------------------


def write_excel(results: List[ItemResult], summary: dict, config: dict, xlsx_path: Path):
    wb = Workbook()
    # Summary sheet
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum["A1"] = "JAMA TRACEABILITY REPORT (Offline DB)"
    ws_sum["A1"].font = Font(bold=True)
    rows = [
        ("Items in scope", summary["total_items"]),
        ("Items evaluated", summary["items_evaluated"]),
        ("Items passed", summary["items_passed"]),
        ("Items failed", summary["items_failed"]),
        ("Checks total", summary["total_checks"]),
        ("Checks passed", summary["checks_passed"]),
        ("Checks failed", summary["checks_failed"]),
    ]
    for r, (k, v) in enumerate(rows, start=3):
        ws_sum[f"A{r}"] = k
        ws_sum[f"B{r}"] = v

    # Results sheet
    ws_res = wb.create_sheet("Results")
    header = [
        "Document Key",
        "Name",
        "Type",
        "Type Description",
        "Overall",
        "Direction",
        "Expression",
        "Check Passed",
        "Found Links",
        "Missing Types",
        "URL",
    ]
    ws_res.append(header)
    for cell in ws_res[1]:
        cell.font = Font(bold=True)
    for r in results:
        base = [
            r.node.document_key,
            r.node.name,
            r.node.type_key or "",
            type_desc(config, r.node.type_key),
            "PASS" if r.overall_passed else "FAIL",
        ]
        if not r.checks:
            ws_res.append(base + ["", "", "", "", r.node.url])
        else:
            for c in r.checks:
                found = "; ".join(f"{n.document_key} ({type_desc(config, n.type_key)})" for n in c.found_nodes)
                missing = "; ".join(type_desc(config, m) for m in c.missing_types)
                ws_res.append(base + [c.direction, c.expression, "PASS" if c.passed else "FAIL", found, missing, r.node.url])

    # Missing Links sheet
    ws_miss = wb.create_sheet("Missing Links")
    ws_miss.append(["Document Key", "Name", "Type", "Direction", "Missing Type", "Missing Type Description", "URL"])
    for cell in ws_miss[1]:
        cell.font = Font(bold=True)

    # Build missing links list from results
    for r in results:
        for c in r.checks:
            if (not c.passed) and c.missing_types:
                for m in c.missing_types:
                    ws_miss.append(
                        [
                            r.node.document_key,
                            r.node.name,
                            r.node.type_key or "",
                            c.direction,
                            m,
                            type_desc(config, m),
                            r.node.url,
                        ]
                    )

    # nice widths
    for ws in (ws_res, ws_miss):
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 36
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 36
        ws.column_dimensions["E"].width = 10
        ws.column_dimensions["F"].width = 14
        ws.column_dimensions["G"].width = 50
        if ws is ws_res:
            ws.column_dimensions["H"].width = 12
            ws.column_dimensions["I"].width = 60
            ws.column_dimensions["J"].width = 36
            ws.column_dimensions["K"].width = 50

    wb.save(xlsx_path)


# --------------------------------- CLI ---------------------------------


def main() -> int:
    ap = argparse.ArgumentParser(description="Analyze traceability from a local Jama graph DB (no live Jama).")
    ap.add_argument("--db-file", required=True, help="Path to JSON DB produced by jamagraph_export.py")
    ap.add_argument("--config", required=True, help="Traceability config JSON (same format as before)")
    ap.add_argument("--output-xlsx", default="traceability_report.xlsx", help="Where to write the Excel report")
    ap.add_argument("--report-format", choices=["text", "json", "csv"], default="text", help="Stdout format")
    ap.add_argument("--keys-file", help="Optional file with document keys to restrict analysis scope (one per line)")
    ap.add_argument("--verbose", action="store_true", help="Extra prints")
    args = ap.parse_args()

    db_path = Path(args.db_file)
    cfg_path = Path(args.config)
    xlsx_path = Path(args.output_xlsx)

    nodes, edges, seeds = load_db(db_path)
    config = load_config(cfg_path)

    # optional scope restriction
    restrict: Optional[Set[str]] = None
    if args.keys_file:
        restrict = {line.strip() for line in Path(args.keys_file).read_text(encoding="utf-8").splitlines() if line.strip()}

    # classify
    classify_nodes(nodes, config)

    # analyze
    results, summary, _missing = analyze_traceability(nodes, edges, config, restrict_doc_keys=restrict)

    # stdout report
    if args.report_format == "text":
        print(render_stdout(results, summary, config))
    else:
        print(render_structured(results, summary, config, args.report_format))

    # excel report
    write_excel(results, summary, config, xlsx_path)
    if args.verbose:
        print(f"[INFO] Excel report written to: {xlsx_path.resolve()}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
